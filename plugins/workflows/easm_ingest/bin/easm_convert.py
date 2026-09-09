#!/usr/bin/env python3
"""
easm_convert.py — ThreatBook EASM deliverables -> "EASM Data Pack" JSON (contract v1.0)

Input (analyst deliverables for ONE reporting period):
  --inventory   Asset inventory workbook (.xls via xlrd, or .xlsx via openpyxl). Lifecycle
                status is read from the cell FILL COLOUR of column A (green=new, orange=closed,
                grey=inactive, cyan=updated, none=active) — the analysts' convention.
  --report      Quarterly EASM report (.docx): Key Findings, per-module narrative,
                recommendations, dark-web / code-leak / file-leak evidence tables + screenshots.
  --verification (optional) Leaked-credential login verification report (.docx): verified
                accounts + login-success screenshots.
  --previous-summary (optional) JSON with the previous period's KPI counts (for delta backfill
                when the previous period has no data pack).

Output (--out DIR):
  DIR/datapack.json                      the contract document
  DIR/assets/screenshots/<module>/*.jpg  downscaled evidence images (referenced from datapack)
  DIR/convert-report.md                  counts, cross-checks against the report, warnings

Secrets: passwords are NEVER written in clear. With EASM_SECRET_KEY (Fernet key, urlsafe
base64) set, `password_enc` (Fernet) and `password_fp` (HMAC-SHA256, keyed) are emitted.
Without it only `password_masked` + `password_length` are emitted.

Runtime deps: openpyxl, xlrd (for .xls), python-docx, pillow, cryptography (only if a key is set).
Designed to run as a Flocks workflow Python node (declare these in metadata.requirements).
"""
from __future__ import annotations

import argparse
import base64
import datetime as dt
import hashlib
import hmac
import io
import json
import os
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path
from typing import Any, Iterable, Optional

SCHEMA_VERSION = "1.0"
GENERATOR_VERSION = "0.1.1"

# ----------------------------------------------------------------------------
# Lifecycle colour convention (RGB of column-A fill)
# ----------------------------------------------------------------------------
COLOUR_LIFECYCLE = {
    (204, 255, 204): "new",       # green  — discovered this period
    (255, 204, 153): "closed",    # orange — closed / inactive / unreachable / removed / expired / taken down
    (192, 192, 192): "inactive",  # grey   — historical, inactive
    (204, 255, 255): "updated",   # cyan   — updated this period (apps)
    (255, 255, 255): "active",    # white  — historical, active
    None: "active",               # no fill — historical, active
}
CURRENT_LIFECYCLES = {"new", "active", "updated"}

# module-specific label for the orange state (what the analysts mean by it per sheet)
CLOSED_LABEL = {
    "domains": "inactive", "ips": "unreachable", "services": "closed", "login_portals": "closed",
    "mobile_apps": "removed", "dark_web": "invalid", "files": "taken_down", "code": "expired",
    "websites": "closed", "components": "closed", "emails": "closed", "credentials": "closed",
}

MODULE_LABELS = OrderedDict([
    ("asset_discovery", "Asset discovery (domains, IPs, websites, services)"),
    ("mobile_apps", "Mobile applications"),
    ("social_accounts", "WeChat official accounts & mini programs"),
    ("login_portals", "Exposed login portals"),
    ("risky_services", "Exposed risky ports & services"),
    ("certificates", "SSL/TLS certificate risks"),
    ("http_misconfigurations", "HTTP security header configuration"),
    ("malicious_ip_tags", "Malicious IP reputation tags"),
    ("vulnerabilities", "Vulnerabilities on internet-facing assets"),
    ("dark_web", "Dark web data leaks"),
    ("files", "Sensitive file leaks (cloud storage)"),
    ("code", "Source code leaks (GitHub/GitLab)"),
    ("credentials", "Leaked login credentials"),
    ("emails", "Exposed corporate email addresses"),
])

# ----------------------------------------------------------------------------
# Small helpers
# ----------------------------------------------------------------------------
def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def stable_id(*parts: Any) -> str:
    raw = "|".join(str(p or "").strip().lower() for p in parts)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]


def clean(v: Any) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, float):
        if v != v:  # NaN
            return None
        if v.is_integer():
            return str(int(v))
        return str(v)
    s = str(v).replace("\xa0", " ").strip()
    return s or None


def parse_list(v: Any) -> list[str]:
    """Parse "['a', 'b']" / "a、b" / "a, b" into a list of strings."""
    s = clean(v)
    if not s:
        return []
    if s.startswith("["):
        items = re.findall(r"'([^']*)'|\"([^\"]*)\"", s)
        out = [a or b for a, b in items]
        if out:
            return [x.strip() for x in out if x.strip()]
        s = s.strip("[]")
    parts = re.split(r"[、,;]\s*", s)
    return [p.strip().strip("'\"") for p in parts if p.strip().strip("'\"")]


def parse_bool_cn(v: Any) -> Optional[bool]:
    s = clean(v)
    if s is None:
        return None
    if s in ("是", "Yes", "yes", "True", "true", "Y"):
        return True
    if s in ("否", "No", "no", "False", "false", "N"):
        return False
    return None


DATE_PATTERNS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M")


def norm_date(v: Any, datemode: int = 0) -> Optional[str]:
    """Return ISO date (YYYY-MM-DD) or ISO datetime string when time is meaningful."""
    if v is None or v == "":
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        if isinstance(v, dt.datetime) and (v.hour or v.minute or v.second):
            return v.strftime("%Y-%m-%d %H:%M:%S")
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and 20000 < v < 80000:  # Excel serial
        try:
            import xlrd  # type: ignore
            d = xlrd.xldate_as_datetime(v, datemode)
            return norm_date(d)
        except Exception:
            return None
    s = clean(v)
    if not s:
        return None
    s2 = s.replace("/", "-")
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?)?", s2)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            if m.group(4):
                return dt.datetime(y, mo, d, int(m.group(4)), int(m.group(5)), int(m.group(6) or 0)).strftime("%Y-%m-%d %H:%M:%S")
            return dt.date(y, mo, d).strftime("%Y-%m-%d")
        except ValueError:
            return s
    m = re.match(r"^(\d{1,2})/([A-Za-z]{3})/(\d{4})$", s)  # 10/Jun/2026
    if m:
        try:
            return dt.datetime.strptime(s, "%d/%b/%Y").strftime("%Y-%m-%d")
        except ValueError:
            return s
    return s


def days_between(iso_date: Optional[str], ref: dt.date) -> Optional[int]:
    if not iso_date:
        return None
    try:
        d = dt.datetime.strptime(iso_date[:10], "%Y-%m-%d").date()
    except ValueError:
        return None
    return (d - ref).days


CITY_COUNTRY = {
    "jakarta": "Indonesia", "singapore": "Singapore", "hong kong": "Hong Kong", "kuala lumpur": "Malaysia",
    "taipei": "Taiwan", "tokyo": "Japan", "mumbai": "India", "london": "United Kingdom", "sydney": "Australia",
    "shanghai": "China", "beijing": "China", "changzhou": "China", "shenzhen": "China", "hangzhou": "China",
    "new york": "United States", "united states new york": "United States", "dubai": "United Arab Emirates",
}


def norm_country(geo: Optional[str]) -> Optional[str]:
    if not geo:
        return None
    g = geo.strip()
    if "," in g:
        return g.split(",")[-1].strip()
    low = g.lower()
    for k, v in CITY_COUNTRY.items():
        if low == k or low.startswith(k + " "):
            return v
    return g


def url_host(u: Optional[str]) -> Optional[str]:
    if not u:
        return None
    m = re.match(r"^(?:[a-z]+://)?([^/:?#]+)", u.strip(), re.I)
    return m.group(1).lower() if m else None


def url_scheme(u: Optional[str]) -> Optional[str]:
    if not u:
        return None
    m = re.match(r"^([a-z]+)://", u.strip(), re.I)
    return m.group(1).lower() if m else None


# ----------------------------------------------------------------------------
# Secrets
# ----------------------------------------------------------------------------
class Secrets:
    def __init__(self, key: Optional[str]):
        self.key = key
        self.fernet = None
        if key:
            from cryptography.fernet import Fernet  # type: ignore
            self.fernet = Fernet(key.encode() if isinstance(key, str) else key)
            self._hmac_key = hashlib.sha256(b"easm-fp:" + (key.encode() if isinstance(key, str) else key)).digest()

    @staticmethod
    def mask(pw: Optional[str]) -> Optional[str]:
        if not pw:
            return None
        keep = 2 if len(pw) > 4 else 1
        return pw[:keep] + "•" * min(8, max(3, len(pw) - keep))

    def encrypt(self, pw: Optional[str]) -> Optional[str]:
        if not pw or not self.fernet:
            return None
        return self.fernet.encrypt(pw.encode("utf-8")).decode("ascii")

    def fingerprint(self, pw: Optional[str]) -> Optional[str]:
        if not pw or not self.key:
            return None
        return hmac.new(self._hmac_key, pw.encode("utf-8"), hashlib.sha256).hexdigest()[:16]


# ----------------------------------------------------------------------------
# Workbook readers (xls via xlrd with colours, xlsx via openpyxl with fills)
# ----------------------------------------------------------------------------
class SheetRows:
    """Uniform access: header, rows (list of cell values), lifecycle per row."""

    def __init__(self, name: str, header: list[Optional[str]], rows: list[list[Any]], lifecycles: list[str]):
        self.name = name
        self.header = header
        self.rows = rows
        self.lifecycles = lifecycles

    def col(self, *names: str) -> Optional[int]:
        norm = [(h or "").replace("\xa0", " ").strip().lower() for h in self.header]
        for n in names:
            n = n.lower()
            for i, h in enumerate(norm):
                if h == n:
                    return i
        for n in names:
            n = n.lower()
            for i, h in enumerate(norm):
                if h and (h.startswith(n) or n in h):
                    return i
        return None


def read_workbook(path: Path) -> tuple[dict[str, SheetRows], int]:
    if path.suffix.lower() == ".xls":
        return _read_xls(path)
    return _read_xlsx(path)


def _lifecycle_from_rgb(rgb: Optional[tuple[int, int, int]]) -> str:
    if rgb is None:
        return "active"
    return COLOUR_LIFECYCLE.get(tuple(rgb), "active")


def _read_xls(path: Path) -> tuple[dict[str, SheetRows], int]:
    import xlrd  # type: ignore
    book = xlrd.open_workbook(str(path), formatting_info=True)
    out: dict[str, SheetRows] = {}
    for sh in book.sheets():
        if sh.nrows == 0:
            continue
        header = [clean(sh.cell_value(0, c)) for c in range(sh.ncols)]
        rows, lifes = [], []
        for r in range(1, sh.nrows):
            vals = []
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if sh.cell_type(r, c) == xlrd.XL_CELL_DATE:
                    v = xlrd.xldate_as_datetime(v, book.datemode)
                vals.append(v)
            if all(clean(v) is None for v in vals):
                continue
            xf = book.xf_list[sh.cell_xf_index(r, 0)]
            rgb = book.colour_map.get(xf.background.pattern_colour_index)
            rows.append(vals)
            lifes.append(_lifecycle_from_rgb(rgb))
        out[sh.name] = SheetRows(sh.name, header, rows, lifes)
    return out, book.datemode


def _read_xlsx(path: Path) -> tuple[dict[str, SheetRows], int]:
    import openpyxl  # type: ignore
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out: dict[str, SheetRows] = {}
    for ws in wb.worksheets:
        header = None
        rows, lifes = [], []
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            vals = [c.value for c in row]
            if header is None:
                header = [clean(v) for v in vals]
                continue
            if all(clean(v) is None for v in vals):
                continue
            c0 = row[0]
            rgb = None
            if c0.fill is not None and c0.fill.fill_type and c0.fill.fgColor is not None and c0.fill.fgColor.type == "rgb":
                argb = c0.fill.fgColor.rgb
                if isinstance(argb, str) and len(argb) >= 6:
                    h = argb[-6:]
                    rgb = (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
            rows.append(vals)
            lifes.append(_lifecycle_from_rgb(rgb))
        if header is not None:
            out[ws.title] = SheetRows(ws.title, header, rows, lifes)
    return out, 0


def find_sheet(sheets: dict[str, SheetRows], *candidates: str) -> Optional[SheetRows]:
    keys = {k.lower().strip(): k for k in sheets}
    for c in candidates:
        if c.lower() in keys:
            return sheets[keys[c.lower()]]
    for c in candidates:
        for k in keys:
            if c.lower() in k:
                return sheets[keys[k]]
    return None


# ----------------------------------------------------------------------------
# Inventory -> entities
# ----------------------------------------------------------------------------
class Inventory:
    def __init__(self, sheets: dict[str, SheetRows], datemode: int, report_date: dt.date, secrets: Secrets, warnings: list[str]):
        self.sheets = sheets
        self.datemode = datemode
        self.report_date = report_date
        self.secrets = secrets
        self.warnings = warnings

    def _life(self, sheet: SheetRows, i: int, module: str) -> tuple[str, str]:
        life = sheet.lifecycles[i]
        label = CLOSED_LABEL.get(module, "closed") if life == "closed" else life
        return life, label

    # -- assets ---------------------------------------------------------------
    def domains(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Domain")
        if not sh:
            return []
        c_root, c_dom, c_ip, c_notes = sh.col("RootDomain", "Root Domain"), sh.col("Domain"), sh.col("IP"), sh.col("Notes")
        # openpyxl/xlrd both give 'Domain' for col(); ensure we do not pick RootDomain for Domain
        if c_dom == c_root:
            c_dom = next((i for i, h in enumerate(sh.header) if (h or "").strip().lower() == "domain"), c_dom)
        out = []
        for i, r in enumerate(sh.rows):
            dom = clean(r[c_dom]) if c_dom is not None else None
            if not dom:
                continue
            life, label = self._life(sh, i, "domains")
            out.append({
                "id": dom.lower(),
                "domain": dom.lower(),
                "root_domain": (clean(r[c_root]) or "").lower() if c_root is not None else None,
                "ips": parse_list(r[c_ip]) if c_ip is not None else [],
                "lifecycle": life, "lifecycle_label": label,
                "analyst_note": clean(r[c_notes]) if c_notes is not None else None,
            })
        return out

    def ips(self) -> list[dict]:
        sh = find_sheet(self.sheets, "IP")
        if not sh:
            return []
        c_ip, c_geo, c_op, c_attr, c_port, c_notes = sh.col("IP Address"), sh.col("Geographic Location"), sh.col("Operater", "Operator"), sh.col("Attribute"), sh.col("Port"), sh.col("Notes")
        out = []
        for i, r in enumerate(sh.rows):
            ip = clean(r[c_ip]) if c_ip is not None else None
            if not ip:
                continue
            life, label = self._life(sh, i, "ips")
            geo = clean(r[c_geo]) if c_geo is not None else None
            ports = parse_list(r[c_port]) if c_port is not None else []
            out.append({
                "id": ip, "ip": ip,
                "geo_raw": geo, "country": norm_country(geo),
                "operator": clean(r[c_op]) if c_op is not None else None,
                "attributes": parse_list(r[c_attr]) if c_attr is not None else [],
                "ports": ports, "open_port_count": len(ports),
                "lifecycle": life, "lifecycle_label": label,
                "analyst_note": clean(r[c_notes]) if c_notes is not None else None,
            })
        return out

    def websites(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Website")
        if not sh:
            return []
        c_url, c_title, c_svc, c_hdr = sh.col("Website"), sh.col("Title"), sh.col("Service"), sh.col("Header")
        out = []
        for i, r in enumerate(sh.rows):
            url = clean(r[c_url]) if c_url is not None else None
            if not url:
                continue
            life, label = self._life(sh, i, "websites")
            headers = None
            server = None
            if c_hdr is not None and clean(r[c_hdr]):
                try:
                    headers = json.loads(str(r[c_hdr]))
                    if isinstance(headers, dict):
                        server = headers.get("Server") or headers.get("server")
                except Exception:
                    headers = {"_raw": str(r[c_hdr])[:2000]}
            out.append({
                "id": url, "url": url, "host": url_host(url), "scheme": url_scheme(url),
                "title": clean(r[c_title]) if c_title is not None else None,
                "technologies": parse_list(r[c_svc]) if c_svc is not None else [],
                "server_header": clean(server),
                "headers": headers,
                "lifecycle": life, "lifecycle_label": label,
            })
        return out

    def services(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Service")
        if not sh:
            return []
        c_ip, c_svc, c_port = sh.col("IP Address"), sh.col("Service"), sh.col("Port")
        out = []
        for i, r in enumerate(sh.rows):
            ip = clean(r[c_ip]) if c_ip is not None else None
            port = clean(r[c_port]) if c_port is not None else None
            if not ip or not port:
                continue
            life, label = self._life(sh, i, "services")
            out.append({
                "id": f"{ip}:{port}", "ip": ip, "port": int(float(port)) if re.match(r"^\d+(\.0)?$", port) else port,
                "service": clean(r[c_svc]) if c_svc is not None else None,
                "lifecycle": life, "lifecycle_label": label,
            })
        return out

    def components(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Components")
        if not sh:
            return []
        c_site, c_comp, c_type, c_first, c_last = sh.col("Website"), sh.col("Web Components"), sh.col("Web Components Types", "Types"), sh.col("First Seen"), sh.col("Last Seen")
        # "Web Components" may match both name and type columns; resolve precisely
        names = [(h or "").strip().lower() for h in sh.header]
        if "web components" in names:
            c_comp = names.index("web components")
        if "web components types" in names:
            c_type = names.index("web components types")
        out = []
        for i, r in enumerate(sh.rows):
            site = clean(r[c_site]) if c_site is not None else None
            comp = clean(r[c_comp]) if c_comp is not None else None
            if not site or not comp:
                continue
            life, label = self._life(sh, i, "components")
            out.append({
                "id": f"{site}|{comp}", "host": site, "component": comp,
                "component_type": clean(r[c_type]) if c_type is not None else None,
                "first_seen": norm_date(r[c_first], self.datemode) if c_first is not None else None,
                "last_seen": norm_date(r[c_last], self.datemode) if c_last is not None else None,
                "lifecycle": life, "lifecycle_label": label,
            })
        return out

    def certificates(self) -> list[dict]:
        sh = find_sheet(self.sheets, "SSL Certificate", "SSL")
        if not sh:
            return []
        c_host, c_port, c_proto, c_low, c_hb, c_n, c_weak, c_valid = (sh.col("Domain/IP"), sh.col("Port"), sh.col("Protocol Version"), sh.col("Low Version"), sh.col("Heartbleed"), sh.col("Number Of Encryption Suites"), sh.col("Weak Algorithm"), sh.col("Certificate Validity"))
        out = []
        seen: set[str] = set()
        dups = 0
        for i, r in enumerate(sh.rows):
            host = clean(r[c_host]) if c_host is not None else None
            if not host:
                continue
            port = clean(r[c_port]) if c_port is not None else None
            proto = clean(r[c_proto]) if c_proto is not None else None
            key = f"{host}:{port}:{proto}".lower()
            if key in seen:  # the sheet repeats the same endpoint/protocol row several times
                dups += 1
                continue
            seen.add(key)
            valid = norm_date(r[c_valid], self.datemode) if c_valid is not None else None
            dte = days_between(valid, self.report_date)
            bucket = None
            if dte is not None:
                bucket = "expired" if dte < 0 else "lt_30d" if dte < 30 else "lt_90d" if dte < 90 else "ge_90d"
            weak = parse_list(r[c_weak]) if c_weak is not None else []
            out.append({
                "id": f"{host}:{port}:{proto}", "host": host, "port": int(float(port)) if port and re.match(r"^\d+(\.0)?$", port) else port,
                "protocol": proto,
                "low_version": parse_bool_cn(r[c_low]) if c_low is not None else None,
                "heartbleed": parse_bool_cn(r[c_hb]) if c_hb is not None else None,
                "cipher_suite_count": int(float(clean(r[c_n]))) if c_n is not None and clean(r[c_n]) and re.match(r"^\d+(\.0)?$", clean(r[c_n])) else None,
                "weak_algorithms": weak,
                "valid_until": valid, "days_to_expiry": dte, "expiry_bucket": bucket,
                "lifecycle": sh.lifecycles[i],
            })
        if dups:
            self.warnings.append(f"certificates: {dups} duplicate endpoint/protocol rows collapsed")
        return out

    def http_configs(self) -> list[dict]:
        sh = find_sheet(self.sheets, "HTTP Configuration", "HTTP")
        if not sh:
            return []
        c_site, c_time = sh.col("Website"), sh.col("Detection Time")
        check_cols = [(i, h) for i, h in enumerate(sh.header) if h and i not in (c_site, c_time)]
        out = []
        for i, r in enumerate(sh.rows):
            site = clean(r[c_site]) if c_site is not None else None
            if not site:
                continue
            checks = {}
            for ci, h in check_cols:
                v = clean(r[ci])
                checks[h] = (True if v in ("True", "TRUE", "true") else False if v in ("False", "FALSE", "false") else v)
            out.append({
                "id": site, "url": site, "host": url_host(site),
                "detected_at": norm_date(r[c_time], self.datemode) if c_time is not None else None,
                "checks": checks,
                "lifecycle": sh.lifecycles[i],
            })
        return out

    def login_portals(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Login Potals", "Login Portals", "Login")
        if not sh:
            return []
        c_title, c_link, c_vc, c_mfa, c_upd = sh.col("Title"), sh.col("Link"), sh.col("Verification Code", "Verification"), sh.col("MFA"), sh.col("Update Time", "Update")
        out = []
        for i, r in enumerate(sh.rows):
            link = clean(r[c_link]) if c_link is not None else None
            if not link:
                continue
            life, label = self._life(sh, i, "login_portals")
            vc = clean(r[c_vc]) if c_vc is not None else None
            mfa = clean(r[c_mfa]) if c_mfa is not None else None
            out.append({
                "id": link, "url": link, "host": url_host(link), "scheme": url_scheme(link),
                "title": clean(r[c_title]) if c_title is not None else None,
                "captcha": (vc.lower() if vc else "unknown"), "mfa": (mfa.lower() if mfa else "unknown"),
                "last_checked": norm_date(r[c_upd], self.datemode) if c_upd is not None else None,
                "lifecycle": life, "lifecycle_label": label,
            })
        return out

    def mobile_apps(self) -> list[dict]:
        sh = find_sheet(self.sheets, "APP")
        if not sh:
            return []
        c_name, c_ver, c_dev, c_src, c_dl, c_upd, c_rec = sh.col("APP Name"), sh.col("Version"), sh.col("Developers"), sh.col("Source"), sh.col("Download Link"), sh.col("Update Time"), sh.col("Recording Time")
        out = []
        for i, r in enumerate(sh.rows):
            name = clean(r[c_name]) if c_name is not None else None
            if not name:
                continue
            life, label = self._life(sh, i, "mobile_apps")
            src = clean(r[c_src]) if c_src is not None else None
            out.append({
                "id": stable_id(name, src), "name": name, "version": clean(r[c_ver]) if c_ver is not None else None,
                "developer": clean(r[c_dev]) if c_dev is not None else None, "store": src,
                "download_url": clean(r[c_dl]) if c_dl is not None else None,
                "app_updated_at": norm_date(r[c_upd], self.datemode) if c_upd is not None else None,
                "recorded_at": norm_date(r[c_rec], self.datemode) if c_rec is not None else None,
                "official_store": bool(src and src.lower() in ("appstore", "app store", "google play")),
                "lifecycle": life, "lifecycle_label": label,
            })
        return out

    def wechat_accounts(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Wechat Official Accounts", "Official Accounts")
        if not sh:
            return []
        c_name, c_owner, c_id, c_rec = sh.col("Wechat Official Accounts Name", "Name"), sh.col("Owner"), sh.col("Wechat ID"), sh.col("Recording Time")
        out = []
        for i, r in enumerate(sh.rows):
            name = clean(r[c_name]) if c_name is not None else None
            if not name:
                continue
            wid = clean(r[c_id]) if c_id is not None else None
            out.append({"id": wid or stable_id(name), "name": name, "owner": clean(r[c_owner]) if c_owner is not None else None,
                        "wechat_id": wid, "recorded_at": norm_date(r[c_rec], self.datemode) if c_rec is not None else None,
                        "lifecycle": sh.lifecycles[i]})
        return out

    def mini_programs(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Wechat Mini Program", "Mini Program")
        if not sh:
            return []
        c_name, c_desc, c_owner, c_id, c_rec, c_notes = sh.col("Wechat Mini Program Name", "Name"), sh.col("Description"), sh.col("Owner"), sh.col("AppID"), sh.col("Recording Time"), sh.col("Notes")
        out = []
        for i, r in enumerate(sh.rows):
            name = clean(r[c_name]) if c_name is not None else None
            if not name:
                continue
            appid = clean(r[c_id]) if c_id is not None else None
            out.append({"id": appid or stable_id(name), "name": name, "description": clean(r[c_desc]) if c_desc is not None else None,
                        "owner": clean(r[c_owner]) if c_owner is not None else None, "app_id": appid,
                        "recorded_at": norm_date(r[c_rec], self.datemode) if c_rec is not None else None,
                        "analyst_note": clean(r[c_notes]) if c_notes is not None else None, "lifecycle": sh.lifecycles[i]})
        return out

    # -- leaks ----------------------------------------------------------------
    def dark_web(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Dark Web")
        if not sh:
            return []
        c_forum, c_title, c_url, c_rel, c_pub, c_hist, c_card, c_conf, c_notes = (sh.col("Forum"), sh.col("Title"), sh.col("Url"), sh.col("Release Time"), sh.col("Publisher"), sh.col("Historical Data Leak"), sh.col("Card Number"), sh.col("confidence"), sh.col("Notes"))
        out = []
        for i, r in enumerate(sh.rows):
            title = clean(r[c_title]) if c_title is not None else None
            if not title:
                continue
            life, label = self._life(sh, i, "dark_web")
            taken_down = bool(re.search(r"taken\s*down", title, re.I))
            conf = clean(r[c_conf]) if c_conf is not None else None
            out.append({
                "id": stable_id("darkweb", title), "forum": clean(r[c_forum]) if c_forum is not None else None,
                "title": title, "title_clean": re.sub(r"[（(]\s*have been taken down\s*[)）]", "", title, flags=re.I).strip(),
                "url": clean(r[c_url]) if c_url is not None else None,
                "posted_at": norm_date(r[c_rel], self.datemode) if c_rel is not None else None,
                "publisher_note": clean(r[c_pub]) if c_pub is not None else None,
                "historical_leak_match": clean(r[c_hist]) if c_hist is not None else None,
                "card_number_format": clean(r[c_card]) if c_card is not None else None,
                "confidence": {"middle": "medium"}.get((conf or "").lower(), (conf or "").lower() or None),
                "analyst_note": clean(r[c_notes]) if c_notes is not None else None,
                "taken_down": taken_down,
                "lifecycle": life, "lifecycle_label": label,
                "screenshot": None, "confidence_dimensions": None, "evaluation": None,
            })
        return out

    def emails(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Mail")
        if not sh:
            return []
        c_mail, c_src, c_notes = sh.col("E-mail Address", "Email"), sh.col("Source"), sh.col("Notes")
        out = []
        for i, r in enumerate(sh.rows):
            mail = clean(r[c_mail]) if c_mail is not None else None
            if not mail:
                continue
            life, label = self._life(sh, i, "emails")
            src = clean(r[c_src]) if c_src is not None else None
            out.append({"id": mail.lower(), "email": mail.lower(), "domain": mail.lower().split("@")[-1] if "@" in mail else None,
                        "source": src, "source_host": url_host(src) if src else None,
                        "analyst_note": clean(r[c_notes]) if c_notes is not None else None,
                        "lifecycle": life, "lifecycle_label": label})
        return out

    def files(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Files")
        if not sh:
            return []
        c_name, c_url, c_up, c_notes = sh.col("File Name"), sh.col("File Url"), sh.col("Upload time"), sh.col("Notes")
        out = []
        for i, r in enumerate(sh.rows):
            url = clean(r[c_url]) if c_url is not None else None
            name = clean(r[c_name]) if c_name is not None else None
            if not url and not name:
                continue
            life, label = self._life(sh, i, "files")
            note = clean(r[c_notes]) if c_notes is not None else None
            out.append({"id": stable_id("file", url or name), "file_name": name, "url": url, "platform": url_host(url) if url else None,
                        "uploaded_at": norm_date(r[c_up], self.datemode) if c_up is not None else None,
                        "analyst_note": note, "taken_down": life == "closed" or bool(note and re.search(r"taken\s*down", note, re.I)),
                        "lifecycle": life, "lifecycle_label": label, "screenshot": None})
        return out

    def code(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Source Code")
        if not sh:
            return []
        c_url, c_desc, c_sub, c_rel, c_sens, c_notes = sh.col("Url"), sh.col("Describe"), sh.col("Submitter"), sh.col("Release Time"), sh.col("Sensitive Data"), sh.col("Notes")
        out = []
        for i, r in enumerate(sh.rows):
            url_cell = clean(r[c_url]) if c_url is not None else None
            if not url_cell:
                continue
            urls = [u.strip() for u in re.split(r"[\s,;]+", url_cell) if u.strip()]  # one cell may hold several repos
            url = urls[0]
            life, label = self._life(sh, i, "code")
            sens = clean(r[c_sens]) if c_sens is not None else None
            note = clean(r[c_notes]) if c_notes is not None else None
            sens_flag = bool(sens and sens.lower() not in ("none", "-", "no", "n/a"))
            suspected = bool(note and re.search(r"suspect|credential|password|database|secret|token|api key", note, re.I))
            out.append({"id": stable_id("code", url), "url": url, "additional_urls": urls[1:], "platform": url_host(url),
                        "repository": re.sub(r"^https?://[^/]+/", "", url.rstrip("/")),
                        "description": clean(r[c_desc]) if c_desc is not None else None,
                        "submitter": clean(r[c_sub]) if c_sub is not None else None,
                        "published_at": norm_date(r[c_rel], self.datemode) if c_rel is not None else None,
                        "sensitive_data": sens if sens and sens.lower() not in ("none", "-") else None,
                        "sensitive_data_found": sens_flag, "sensitive_data_suspected": suspected and not sens_flag,
                        "analyst_note": note, "lifecycle": life, "lifecycle_label": label,
                        "screenshot": None, "sensitive_screenshot": None})
        return out

    def credentials(self) -> list[dict]:
        sh = find_sheet(self.sheets, "login credential", "credential")
        if not sh:
            return []
        c_url, c_user, c_pw = sh.col("URL"), sh.col("username"), sh.col("password")
        out = []
        for i, r in enumerate(sh.rows):
            url = clean(r[c_url]) if c_url is not None else None
            user = clean(r[c_user]) if c_user is not None else None
            pw = clean(r[c_pw]) if c_pw is not None else None
            if not user:
                continue
            life = sh.lifecycles[i]
            rec = {
                "id": stable_id("cred", url, user, Secrets.mask(pw), len(pw) if pw else 0), "login_url": url, "host": url_host(url), "username": user,
                "username_type": "email" if user and "@" in user else "account",
                "password_masked": Secrets.mask(pw), "password_length": len(pw) if pw else None,
                "password_enc": self.secrets.encrypt(pw), "password_fp": self.secrets.fingerprint(pw),
                "verified_login": life == "new",  # analysts mark verified accounts green in this sheet
                "verification": None,
                "lifecycle": "active", "lifecycle_label": "active",
            }
            out.append(rec)
        return out

    def vulnerabilities(self) -> list[dict]:
        sh = find_sheet(self.sheets, "Risk", "Vulnerab")
        if not sh:
            return []
        c_v, c_d, c_lvl, c_url, c_ver, c_fix = sh.col("Vulnerability"), sh.col("Describe"), sh.col("Threat level"), sh.col("URL"), sh.col("Verify"), sh.col("Repair")
        out = []
        for i, r in enumerate(sh.rows):
            v = clean(r[c_v]) if c_v is not None else None
            if not v:
                continue
            out.append({"id": stable_id("vuln", v, clean(r[c_url]) if c_url is not None else None), "name": v,
                        "description": clean(r[c_d]) if c_d is not None else None, "threat_level": clean(r[c_lvl]) if c_lvl is not None else None,
                        "url": clean(r[c_url]) if c_url is not None else None, "verified": clean(r[c_ver]) if c_ver is not None else None,
                        "remediation": clean(r[c_fix]) if c_fix is not None else None, "lifecycle": sh.lifecycles[i]})
        return out


# ----------------------------------------------------------------------------
# Report (.docx) parsing
# ----------------------------------------------------------------------------
MODULE_KEYWORDS = [
    ("login_portals", r"login\s*p\w*tal"),  # tolerates the 'Potals' typo in the analysts' template
    ("risky_services", r"risk\s*(ports?|services?)|ports? and services"),
    ("certificates", r"ssl|certificate"),
    ("http_misconfigurations", r"http\s*non-?compliance|http\s*config|http\s*header"),
    ("malicious_ip_tags", r"malicious(ly)?\s*(ip|flag)|ip\s*tags?"),
    ("vulnerabilities", r"vulnerab"),
    ("dark_web", r"dark\s*web"),
    ("files", r"file\s*leak|cloud\s*storage|openly\s*shared"),
    ("code", r"code\s*leak|github|gitlab|code\s*hosting"),
    ("credentials", r"credential"),
    ("emails", r"e-?mail"),
    ("mobile_apps", r"android|apk|\bapps?\b"),
    ("social_accounts", r"wechat|mini\s*program|official\s*account"),
    ("asset_discovery", r"subdomain|ip assets|websites|asset inventory"),
]


def module_for(text: str) -> Optional[str]:
    t = text.lower()
    for mod, pat in MODULE_KEYWORDS:
        if re.search(pat, t):
            return mod
    return None


class DocxReport:
    def __init__(self, path: Path, warnings: list[str]):
        import docx  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        self.docx = docx
        self.qn = qn
        self.doc = docx.Document(str(path))
        self.warnings = warnings
        self.blocks = self._blocks()

    def _blocks(self) -> list[dict]:
        """Linear list of {kind: 'p'|'tbl', heading:(level,text), ...} with the enclosing heading."""
        d = self.doc
        body = d.element.body
        out = []
        cur = None
        for child in body.iterchildren():
            tag = child.tag.split("}")[1]
            if tag == "p":
                p = self.docx.text.paragraph.Paragraph(child, d)
                style = p.style.name if p.style is not None else ""
                text = p.text.replace("\xa0", " ").strip()
                m = re.match(r"Heading (\d)", style)
                if m and text:
                    cur = (int(m.group(1)), text)
                    out.append({"kind": "h", "level": int(m.group(1)), "text": text, "heading": cur})
                    continue
                imgs = self._images_in(child)
                if text or imgs:
                    out.append({"kind": "p", "text": text, "style": style, "images": imgs, "heading": cur})
            elif tag == "tbl":
                tbl = self.docx.table.Table(child, d)
                rows = []
                for row in tbl.rows:
                    cells, seen = [], set()
                    for c in row.cells:
                        if id(c._tc) in seen:
                            continue
                        seen.add(id(c._tc))
                        cells.append({"text": c.text.replace("\xa0", " ").strip(), "images": self._images_in(c._tc)})
                    rows.append(cells)
                out.append({"kind": "tbl", "rows": rows, "heading": cur})
        return out

    def _images_in(self, el) -> list[bytes]:
        blobs = []
        for blip in el.findall(".//" + self.qn("a:blip")):
            rid = blip.get(self.qn("r:embed"))
            if not rid:
                continue
            try:
                part = self.doc.part.related_parts[rid]
                blobs.append(part.blob)
            except KeyError:
                continue
        return blobs

    # -- sections ---------------------------------------------------------------
    def key_findings(self) -> tuple[list[dict], list[str]]:
        items, scope_lines = [], []
        in_kf = False
        for b in self.blocks:
            if b["kind"] == "h":
                in_kf = bool(re.search(r"key\s*findings", b["text"], re.I))
                continue
            if not in_kf or b["kind"] != "p" or not b["text"]:
                continue
            t = b["text"]
            m = re.match(r"^[【\[]\s*([HML])\s*[】\]]\s*(.*)$", t)
            if m:
                sev = {"H": "high", "M": "medium", "L": "low"}[m.group(1)]
                body = m.group(2).strip().rstrip("；;。.").strip()
                items.append({"severity": sev, "module": module_for(body), "text": body,
                              "count": _first_int(body)})
            elif re.match(r"^(no |there (are|is) no|none)", t, re.I) or re.search(r"^no [a-z ]+ (found|identified|expose|are maliciously)", t, re.I) or re.search(r"^no vulnerabilit", t, re.I):
                items.append({"severity": "none", "module": module_for(t), "text": t.rstrip("；;。.").strip(), "count": 0})
            elif re.match(r"^[A-Za-z ]+:\s*$", t):
                continue  # bare group label such as "Vulnerability Risks:" — carries no information
            else:
                scope_lines.append(t)
        return items, scope_lines

    def sections(self) -> dict[str, dict]:
        """module -> {heading, count_in_heading, risk_statement, notes[], tables:[...]}"""
        out: dict[str, dict] = {}
        for b in self.blocks:
            h = b["heading"]
            if not h:
                continue
            level, text = h
            if re.search(r"recommendation", text, re.I) or re.search(r"key\s*findings", text, re.I):
                continue
            mod = module_for(text)
            if not mod:
                continue
            sec = out.setdefault(mod, {"heading": text, "count_in_heading": _count_in_heading(text), "risk_statement": None, "notes": [], "tables": []})
            if b["kind"] == "p" and b["text"]:
                t = b["text"]
                if re.match(r"^risk\s*[:：]", t, re.I):
                    sec["risk_statement"] = re.sub(r"^risk\s*[:：]\s*", "", t, flags=re.I)
                elif not re.search(r"the table includes only the top", t, re.I):
                    sec["notes"].append(t)
            elif b["kind"] == "tbl":
                sec["tables"].append(b["rows"])
        return out

    def recommendations(self) -> list[dict]:
        out = []
        cur = None
        for b in self.blocks:
            if b["kind"] == "h":
                if re.search(r"recommendations?\s+for", b["text"], re.I):
                    mod = module_for(re.sub(r"^.*recommendations?\s+for\s*", "", b["text"], flags=re.I))
                    cur = {"module": mod, "title": re.sub(r"^\d+(\.\d+)*\s*", "", b["text"]).strip(), "items": []}
                    out.append(cur)
                elif b["level"] <= 2 and not re.search(r"recommendation", b["text"], re.I):
                    cur = None
                continue
            if cur and b["kind"] == "p" and b["text"]:
                cur["items"].append(b["text"])
        return out

    # -- evidence tables --------------------------------------------------------
    def index_tables(self) -> list[dict]:
        """Vertical key/value tables that start with an 'Index' row (dark web & code leak items)."""
        out = []
        for b in self.blocks:
            if b["kind"] != "tbl":
                continue
            rows = b["rows"]
            if not rows or not rows[0] or rows[0][0]["text"].strip().lower() != "index":
                continue
            kv, images = {}, []
            for r in rows:
                if not r:
                    continue
                k = r[0]["text"].strip()
                v = r[1]["text"].strip() if len(r) > 1 else ""
                imgs = [im for c in r for im in c["images"]]
                if imgs:
                    images.append((k, imgs))
                kv[k] = v
            out.append({"kv": kv, "images": images, "heading": b["heading"]})
        return out

    def file_leak_table(self) -> list[dict]:
        for b in self.blocks:
            if b["kind"] != "tbl":
                continue
            rows = b["rows"]
            if rows and rows[0] and rows[0][0]["text"].strip().lower() == "file name":
                hdr = [c["text"].strip().lower() for c in rows[0]]
                items = []
                for r in rows[1:]:
                    rec = {hdr[i]: r[i]["text"] for i in range(min(len(hdr), len(r)))}
                    rec["_images"] = [im for c in r for im in c["images"]]
                    items.append(rec)
                return items
        return []

    def credential_verification_records(self) -> list[dict]:
        """For the 'Login Verification' report: one table per record."""
        out = []
        date = None
        for b in self.blocks:
            if b["kind"] == "p" and re.match(r"^date\s*:", b["text"], re.I):
                date = norm_date(re.sub(r"^date\s*:\s*", "", b["text"], flags=re.I))
        for b in self.blocks:
            if b["kind"] != "tbl":
                continue
            rows = b["rows"]
            if not rows or not re.search(r"credential leak record", rows[0][0]["text"], re.I):
                continue
            kv, imgs = {}, []
            for r in rows:
                k = r[0]["text"].strip()
                v = r[1]["text"].strip() if len(r) > 1 else ""
                kv[k] = v
                imgs += [im for c in r for im in c["images"]]
            out.append({
                "record": rows[0][0]["text"].strip(),
                "username": kv.get("Leaked Account / Username"),
                "login_url": kv.get("Login URL"),
                "result": kv.get("Verification Result"),
                "verified_at": date,
                "_images": imgs,
            })
        return out


def _first_int(text: str) -> Optional[int]:
    m = re.search(r"(\d[\d,]*)", text)
    return int(m.group(1).replace(",", "")) if m else None


def _count_in_heading(text: str) -> Optional[int]:
    m = re.search(r"\(\s*(\d+)\s*items?\s*\)", text, re.I)
    if m:
        return int(m.group(1))
    if re.search(r"\(\s*none\s*\)", text, re.I):
        return 0
    return None


# ----------------------------------------------------------------------------
# Images
# ----------------------------------------------------------------------------
class ImageStore:
    def __init__(self, out_dir: Path, max_width: int = 1200, quality: int = 80):
        self.root = out_dir / "assets" / "screenshots"
        self.root.mkdir(parents=True, exist_ok=True)
        self.max_width = max_width
        self.quality = quality
        self.count = 0
        self.bytes_in = 0
        self.bytes_out = 0

    def save(self, module: str, key: str, blob: bytes, suffix: str = "") -> Optional[dict]:
        try:
            from PIL import Image  # type: ignore
        except ImportError:
            return None
        try:
            im = Image.open(io.BytesIO(blob))
            im.load()
        except Exception:
            return None
        self.bytes_in += len(blob)
        w, h = im.size
        if w > self.max_width:
            nh = int(h * self.max_width / w)
            im = im.resize((self.max_width, nh))
        if im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        d = self.root / module
        d.mkdir(parents=True, exist_ok=True)
        name = re.sub(r"[^A-Za-z0-9_.-]", "_", key)[:80] + (f"-{suffix}" if suffix else "") + ".jpg"
        path = d / name
        im.save(path, "JPEG", quality=self.quality, optimize=True)
        size = path.stat().st_size
        self.bytes_out += size
        self.count += 1
        return {"path": str(path.relative_to(self.root.parent.parent)).replace("\\", "/"), "width": im.size[0], "height": im.size[1], "bytes": size,
                "sensitivity": "restricted"}


# ----------------------------------------------------------------------------
# Rules (severity + exposure index) — rules-v1, no AI involvement
# ----------------------------------------------------------------------------
SEVERITY_RANK = {"high": 3, "medium": 2, "low": 1, "info": 0, "none": -1}

EXPOSURE_WEIGHTS = OrderedDict([
    # module: (max points, saturation threshold in weighted units)
    ("credentials", (25, 60)),
    ("dark_web", (20, 6)),
    ("code", (10, 10)),
    ("files", (5, 5)),
    ("login_portals", (10, 300)),
    ("certificates", (10, 10)),
    ("vulnerabilities", (10, 5)),
    ("risky_services", (5, 10)),
    ("malicious_ip_tags", (5, 5)),
])


def severity_login_portal(p: dict) -> str:
    if p.get("scheme") == "http":
        return "medium"
    return "low"


def severity_certificate(c: dict) -> Optional[str]:
    """Only analyst-comparable risk: legacy protocol. Expiry/weak-suite hygiene is derived & informational."""
    proto = (c.get("protocol") or "").upper()
    if proto in ("TLSV1.0", "TLSV1.1", "SSLV3", "SSLV2", "TLSV1"):
        return "medium"
    return None


def severity_code(c: dict) -> str:
    if c.get("sensitive_data_found"):
        return "high"
    if c.get("sensitive_data_suspected"):
        return "medium"
    return "low"


def severity_dark_web(_d: dict) -> str:
    return "high"


def severity_file(_f: dict) -> str:
    return "medium"


def severity_credential(c: dict) -> str:
    return "high" if c.get("verified_login") else "medium"


def severity_vulnerability(v: dict) -> str:
    lvl = (v.get("threat_level") or "").lower()
    if lvl in ("critical", "high", "严重", "高危"):
        return "high"
    if lvl in ("medium", "中危"):
        return "medium"
    return "low"


def exposure_index(module_units: dict[str, float]) -> dict:
    """Saturating log curve per module: points = max * min(1, ln(1+u)/ln(1+threshold))."""
    import math
    mods, total = [], 0.0
    for mod, (mx, thr) in EXPOSURE_WEIGHTS.items():
        u = float(module_units.get(mod, 0.0))
        pts = mx * min(1.0, math.log1p(max(u, 0.0)) / math.log1p(thr)) if u > 0 else 0.0
        pts = round(pts, 1)
        total += pts
        mods.append({"module": mod, "label": MODULE_LABELS.get(mod, mod), "units": round(u, 1), "points": pts, "max_points": mx, "saturation_threshold": thr})
    score = round(total, 1)
    band = "critical" if score >= 70 else "high" if score >= 50 else "medium" if score >= 30 else "low"
    grade = "E" if score >= 70 else "D" if score >= 50 else "C" if score >= 30 else "B" if score >= 15 else "A"
    return {"method": "rules-v1", "score": score, "max_score": 100, "band": band, "grade": grade,
            "direction": "higher = more exposure", "modules": mods,
            "notes": ["Points per module saturate on a log curve at the module's threshold; weights sum to 100.",
                      "Credential units = verified accounts x3 + other accounts x1. Dark web units = open items x1 (taken-down items excluded).",
                      "Code units = high x3 + medium x1.5 + low x0.5. Login portals = current portals x1 + plain-HTTP portals x1 extra.",
                      "Certificate units = analyst-flagged legacy-protocol endpoints x1. HTTP header findings carry no weight until field semantics are confirmed."]}


# ----------------------------------------------------------------------------
# Assembly
# ----------------------------------------------------------------------------
def by_lifecycle(items: Iterable[dict]) -> dict[str, int]:
    c = Counter(i.get("lifecycle", "active") for i in items)
    return {k: c.get(k, 0) for k in ("new", "active", "updated", "closed", "inactive")}


def current(items: Iterable[dict]) -> list[dict]:
    return [i for i in items if i.get("lifecycle", "active") in CURRENT_LIFECYCLES]


def dedupe_ids(items: list[dict], warnings: Optional[list[str]] = None, label: str = "") -> dict:
    """Collapse rows that share an `id` (same login URL, same ip:port, same app+store...).

    Analyst sheets occasionally list one entity twice with a different title / service name /
    download link. The portal shows unique entities, so the first row is kept, its empty fields
    are filled from later rows, a current lifecycle wins over closed/inactive, differing titles
    are kept in `alt_titles`, and `duplicates_collapsed` records how many rows were merged.
    Returns {"rows": removed rows, "current": removed rows that had a current lifecycle} so the
    cross-check against the analyst report (which counts the duplicate rows) can be adjusted.
    """
    by_id: dict[Any, dict] = {}
    kept: list[dict] = []
    removed = removed_current = 0
    for rec in items:
        rid = rec.get("id")
        if rid is None or rid not in by_id:
            if rid is not None:
                by_id[rid] = rec
            kept.append(rec)
            continue
        base = by_id[rid]
        removed += 1
        if rec.get("lifecycle", "active") in CURRENT_LIFECYCLES:
            removed_current += 1
        for k, v in rec.items():
            if k == "id":
                continue
            if base.get(k) in (None, "", [], {}) and v not in (None, "", [], {}):
                base[k] = v
            elif k == "title" and v and base.get(k) and v != base[k]:
                alts = base.setdefault("alt_titles", [])
                if v not in alts:
                    alts.append(v)
        if base.get("lifecycle", "active") not in CURRENT_LIFECYCLES and rec.get("lifecycle", "active") in CURRENT_LIFECYCLES:
            base["lifecycle"] = rec["lifecycle"]
            base["lifecycle_label"] = rec.get("lifecycle_label", base.get("lifecycle_label"))
        base["duplicates_collapsed"] = int(base.get("duplicates_collapsed") or 0) + 1
    items[:] = kept
    if removed and warnings is not None:
        warnings.append(f"{label or 'module'}: {removed} duplicate row(s) collapsed into existing entities (same id)")
    return {"rows": removed, "current": removed_current}


def build(args) -> dict:
    warnings: list[str] = []
    report_date = dt.datetime.strptime(args.report_date, "%Y-%m-%d").date()
    secrets = Secrets(os.environ.get(args.secret_key_env) or None)
    if not secrets.key:
        warnings.append(f"{args.secret_key_env} not set: passwords emitted masked only (no password_enc / password_fp)")

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    images = ImageStore(out_dir, max_width=args.max_image_width)

    sources = []
    sheets, datemode = read_workbook(Path(args.inventory))
    sources.append({"kind": "inventory", "file": Path(args.inventory).name, "sha256": sha256_file(Path(args.inventory)), "sheets": list(sheets.keys())})
    inv = Inventory(sheets, datemode, report_date, secrets, warnings)

    assets = OrderedDict([
        ("domains", inv.domains()), ("ips", inv.ips()), ("websites", inv.websites()), ("services", inv.services()),
        ("components", inv.components()), ("certificates", inv.certificates()), ("http_configs", inv.http_configs()),
        ("mobile_apps", inv.mobile_apps()), ("wechat_official_accounts", inv.wechat_accounts()), ("wechat_mini_programs", inv.mini_programs()),
    ])
    login_portals = inv.login_portals()
    vulnerabilities = inv.vulnerabilities()
    leaks = OrderedDict([("dark_web", inv.dark_web()), ("files", inv.files()), ("code", inv.code()), ("credentials", inv.credentials()), ("emails", inv.emails())])

    # -- entity ids must be unique per module: collapse duplicate rows before any KPI is computed --
    collapsed: "OrderedDict[str, dict]" = OrderedDict()
    for _name, _lst in [*assets.items(), ("login_portals", login_portals), ("vulnerabilities", vulnerabilities), *leaks.items()]:
        _r = dedupe_ids(_lst, warnings, _name)
        if _r["rows"]:
            collapsed[_name] = _r

    # -- report narrative + evidence -------------------------------------------
    narrative = {"key_findings": [], "scope_text": [], "modules": {}, "recommendations": []}
    section_counts = {}
    if args.report:
        rep = DocxReport(Path(args.report), warnings)
        sources.append({"kind": "report", "file": Path(args.report).name, "sha256": sha256_file(Path(args.report))})
        kf, scope = rep.key_findings()
        narrative["key_findings"] = kf
        narrative["scope_text"] = scope
        secs = rep.sections()
        for mod, sec in secs.items():
            narrative["modules"][mod] = {"heading": sec["heading"], "count_in_heading": sec["count_in_heading"], "risk_statement": sec["risk_statement"], "period_notes": sec["notes"]}
            if sec["count_in_heading"] is not None:
                section_counts[mod] = sec["count_in_heading"]
        narrative["recommendations"] = rep.recommendations()

        # dark web + code leak evidence
        dw_by_title = {d["title_clean"].lower(): d for d in leaks["dark_web"]}
        dw_by_titlefull = {d["title"].lower(): d for d in leaks["dark_web"]}
        code_by_url = {}
        for c in leaks["code"]:
            for u in [c["url"]] + c.get("additional_urls", []):
                code_by_url[u.rstrip("/").lower()] = c
        for it in rep.index_tables():
            kv, imgs = it["kv"], it["images"]
            if "Dark Web Links" in kv or "Website Title" in kv:
                title = (kv.get("Website Title") or "").strip()
                tclean = re.sub(r"[（(]\s*have been taken down\s*[)）]", "", title, flags=re.I).strip().lower()
                d = dw_by_title.get(tclean) or dw_by_titlefull.get(title.lower())
                if not d:
                    # fuzzy: prefix match
                    d = next((x for k, x in dw_by_title.items() if k[:25] and tclean.startswith(k[:25])), None)
                if not d:
                    warnings.append(f"dark web evidence not matched to inventory: {title[:60]}")
                    continue
                d["confidence_overall"] = kv.get("Overall Confidence Level")
                d["confidence_dimensions"] = kv.get("Scoring Dimensions of Confidence Level")
                d["evaluation"] = kv.get("Evaluation Details")
                d["posted_at"] = d["posted_at"] or norm_date(kv.get("Release Time"))
                if not d.get("taken_down") and re.search(r"taken\s*down", title, re.I):
                    d["taken_down"] = True
                shots = [im for k, ims in imgs for im in ims]
                if shots:
                    d["screenshot"] = images.save("dark_web", d["id"], shots[0])
                    if len(shots) > 1:
                        d["screenshots_extra"] = [images.save("dark_web", d["id"], s, suffix=str(n + 2)) for n, s in enumerate(shots[1:])]
            elif "Url" in kv:
                url = (kv.get("Url") or "").rstrip("/").lower()
                c = code_by_url.get(url)
                if not c:
                    warnings.append(f"code leak evidence not matched to inventory: {url[:60]}")
                    continue
                c["description"] = c["description"] or (kv.get("Describe") if kv.get("Describe") not in (None, "-", "") else None)
                c["analyst_note"] = c["analyst_note"] or kv.get("Notes")
                if kv.get("Sensitive Data") and kv.get("Sensitive Data").lower() not in ("none", "-", ""):
                    c["sensitive_data"] = c["sensitive_data"] or kv.get("Sensitive Data")
                    c["sensitive_data_found"] = True
                for k, ims in imgs:
                    if re.search(r"sensitive", k, re.I):
                        c["sensitive_screenshot"] = images.save("code", c["id"], ims[0], suffix="sensitive")
                        if not c["sensitive_data_found"]:
                            c["sensitive_data_suspected"] = True
                    elif c.get("screenshot") is None:
                        c["screenshot"] = images.save("code", c["id"], ims[0])
        # file leak screenshots
        files_by_url = {f["url"]: f for f in leaks["files"] if f.get("url")}
        for row in rep.file_leak_table():
            f = files_by_url.get(row.get("file url"))
            if not f:
                continue
            if row.get("_images"):
                f["screenshot"] = images.save("files", f["id"], row["_images"][0])
            if not f.get("analyst_note") and row.get("notes"):
                f["analyst_note"] = row["notes"]

    # -- verification report ----------------------------------------------------
    if args.verification:
        ver = DocxReport(Path(args.verification), warnings)
        sources.append({"kind": "verification", "file": Path(args.verification).name, "sha256": sha256_file(Path(args.verification))})
        creds_by_user = defaultdict(list)
        for c in leaks["credentials"]:
            creds_by_user[(c["username"] or "").lower()].append(c)
        matched = 0
        for rec in ver.credential_verification_records():
            cands = creds_by_user.get((rec["username"] or "").lower(), [])
            if not cands:
                warnings.append(f"verified credential not found in inventory: {rec['username']}")
                continue
            c = next((x for x in cands if (x.get("login_url") or "").rstrip("/") == (rec.get("login_url") or "").rstrip("/")), cands[0])
            c["verified_login"] = True
            c["verification"] = {"record": rec["record"], "result": rec["result"], "verified_at": rec["verified_at"], "login_url": rec["login_url"],
                                 "screenshot": images.save("credentials", c["id"], rec["_images"][0]) if rec["_images"] else None}
            matched += 1
        if matched:
            narrative["modules"].setdefault("credentials", {})["verified_records"] = matched

    # -- risks ----------------------------------------------------------------------
    for p in login_portals:
        p["severity"] = severity_login_portal(p)
    cert_risks, cert_hygiene = [], {"expired": 0, "lt_30d": 0, "lt_90d": 0, "ge_90d": 0, "unknown": 0, "legacy_protocol": 0, "weak_suite_listed": 0, "heartbleed": 0}
    legacy_by_endpoint: "OrderedDict[str, dict]" = OrderedDict()
    for c in assets["certificates"]:
        sev = severity_certificate(c)
        cert_hygiene[c.get("expiry_bucket") or "unknown"] += 1
        if c.get("weak_algorithms"):
            cert_hygiene["weak_suite_listed"] += 1
        if c.get("heartbleed"):
            cert_hygiene["heartbleed"] += 1
        if sev:
            cert_hygiene["legacy_protocol"] += 1
            ep = f"{c['host']}:{c['port']}"
            risk = legacy_by_endpoint.get(ep)
            if risk is None:
                risk = {"id": ep, "host": c["host"], "port": c["port"], "protocols": [], "valid_until": c.get("valid_until"),
                        "days_to_expiry": c.get("days_to_expiry"), "expiry_bucket": c.get("expiry_bucket"), "lifecycle": "active",
                        "severity": sev, "reason": None}
                legacy_by_endpoint[ep] = risk
            if c.get("protocol") and c["protocol"] not in risk["protocols"]:
                risk["protocols"].append(c["protocol"])
            risk["reason"] = "Legacy protocol(s) still accepted: " + ", ".join(risk["protocols"])
    cert_risks = list(legacy_by_endpoint.values())  # one risk per endpoint (matches the analysts' per-host reporting)
    # http header checks (semantics pending): count per header of "issue present" using the analysts' convention (True = observed)
    http_summary = Counter()
    for h in assets["http_configs"]:
        for k, v in h["checks"].items():
            if v is True:
                http_summary[k] += 1
    open_ports = Counter()
    for s in current(assets["services"]):
        open_ports[str(s["port"])] += 1

    risks = OrderedDict([
        ("login_portals", login_portals),
        ("risky_services", []),                     # analyst-flagged risky services (none this period) — open-port inventory lives in assets.services
        ("certificates", cert_risks),
        ("http_misconfigurations", {"semantics_pending": True, "sites_checked": len(assets["http_configs"]), "issue_counts_by_header": dict(http_summary)}),
        ("malicious_ip_tags", []),
        ("vulnerabilities", vulnerabilities),
    ])
    for v in risks["vulnerabilities"]:
        v["severity"] = severity_vulnerability(v)
    for d in leaks["dark_web"]:
        d["severity"] = severity_dark_web(d)
    for f in leaks["files"]:
        f["severity"] = severity_file(f)
    for c in leaks["code"]:
        c["severity"] = severity_code(c)
    for c in leaks["credentials"]:
        c["severity"] = severity_credential(c)
    for e in leaks["emails"]:
        e["severity"] = "low"

    # -- summary -----------------------------------------------------------------------
    cur_portals = current(login_portals)
    dw_open = [d for d in current(leaks["dark_web"]) if not d.get("taken_down")]
    summary = OrderedDict([
        ("assets", OrderedDict([
            ("domains", {"total_rows": len(assets["domains"]), "current": len(current(assets["domains"])), "by_lifecycle": by_lifecycle(assets["domains"]), "root_domains": len({d["root_domain"] for d in assets["domains"] if d.get("root_domain")})}),
            ("ips", {"total_rows": len(assets["ips"]), "current": len(current(assets["ips"])), "by_lifecycle": by_lifecycle(assets["ips"]),
                     "with_open_ports": len([i for i in assets["ips"] if i["open_port_count"]]), "by_country": dict(Counter(i["country"] or "Unknown" for i in assets["ips"]).most_common(12))}),
            ("websites", {"total_rows": len(assets["websites"]), "current": len(current(assets["websites"])), "by_lifecycle": by_lifecycle(assets["websites"])}),
            ("services", {"total_rows": len(assets["services"]), "current": len(current(assets["services"])), "by_lifecycle": by_lifecycle(assets["services"]),
                          "ips_with_current_services": len({s["ip"] for s in current(assets["services"])}), "top_ports": dict(open_ports.most_common(15))}),
            ("components", {"total_rows": len(assets["components"]), "current": len(current(assets["components"])), "by_lifecycle": by_lifecycle(assets["components"]),
                            "by_type": dict(Counter(c["component_type"] or "unknown" for c in assets["components"]).most_common(10)), "admin_panels": len([c for c in assets["components"] if (c.get("component_type") or "").lower() == "admin panel"])}),
            ("certificates", {"total_rows": len(assets["certificates"]), "hygiene": cert_hygiene, "by_protocol": dict(Counter(c["protocol"] or "unknown" for c in assets["certificates"]))}),
            ("http_configs", {"total_rows": len(assets["http_configs"]), "semantics_pending": True}),
            ("mobile_apps", {"total_rows": len(assets["mobile_apps"]), "current": len(current(assets["mobile_apps"])), "by_lifecycle": by_lifecycle(assets["mobile_apps"]),
                             "by_store": dict(Counter(a["store"] or "unknown" for a in assets["mobile_apps"]).most_common(12)), "official_store": len([a for a in assets["mobile_apps"] if a["official_store"]])}),
            ("wechat_official_accounts", {"total_rows": len(assets["wechat_official_accounts"])}),
            ("wechat_mini_programs", {"total_rows": len(assets["wechat_mini_programs"])}),
        ])),
        ("risks", OrderedDict([
            ("login_portals", {"current": len(cur_portals), "new": len([p for p in login_portals if p["lifecycle"] == "new"]), "closed": len([p for p in login_portals if p["lifecycle"] == "closed"]),
                               "plain_http": len([p for p in cur_portals if p.get("scheme") == "http"]), "mfa": dict(Counter(p["mfa"] for p in cur_portals)), "captcha": dict(Counter(p["captcha"] for p in cur_portals)),
                               "by_severity": dict(Counter(p["severity"] for p in cur_portals))}),
            ("risky_services", {"current": 0}),
            ("certificates", {"current": len(cert_risks), "by_severity": dict(Counter(c["severity"] for c in cert_risks))}),
            ("http_misconfigurations", {"sites_checked": len(assets["http_configs"]), "semantics_pending": True}),
            ("malicious_ip_tags", {"current": 0}),
            ("vulnerabilities", {"current": len(vulnerabilities)}),
        ])),
        ("leaks", OrderedDict([
            ("dark_web", {"total_rows": len(leaks["dark_web"]), "open": len(dw_open), "taken_down": len([d for d in leaks["dark_web"] if d.get("taken_down") or d["lifecycle"] == "closed"]),
                          "new": len([d for d in leaks["dark_web"] if d["lifecycle"] == "new"]), "by_forum": dict(Counter(d["forum"] or "unknown" for d in leaks["dark_web"]).most_common(8)),
                          "with_screenshot": len([d for d in leaks["dark_web"] if d.get("screenshot")])}),
            ("files", {"total_rows": len(leaks["files"]), "open": len([f for f in leaks["files"] if not f["taken_down"]]), "taken_down": len([f for f in leaks["files"] if f["taken_down"]]), "with_screenshot": len([f for f in leaks["files"] if f.get("screenshot")])}),
            ("code", {"total_rows": len(leaks["code"]), "current": len(current(leaks["code"])), "by_lifecycle": by_lifecycle(leaks["code"]), "by_severity": dict(Counter(c["severity"] for c in current(leaks["code"]))),
                      "sensitive_found": len([c for c in leaks["code"] if c["sensitive_data_found"]]), "with_screenshot": len([c for c in leaks["code"] if c.get("screenshot")])}),
            ("credentials", {"total": len(leaks["credentials"]), "verified_login": len([c for c in leaks["credentials"] if c["verified_login"]]), "by_host": dict(Counter(c["host"] or "unknown" for c in leaks["credentials"]).most_common(10)),
                             "by_severity": dict(Counter(c["severity"] for c in leaks["credentials"])), "with_verification_screenshot": len([c for c in leaks["credentials"] if c.get("verification") and c["verification"].get("screenshot")]),
                             "encrypted": bool(secrets.key)}),
            ("emails", {"total": len(leaks["emails"]), "new": len([e for e in leaks["emails"] if e["lifecycle"] == "new"]), "by_source": dict(Counter(e["source_host"] or e["source"] or "unknown" for e in leaks["emails"]).most_common(8))}),
        ])),
    ])

    summary["duplicates_collapsed"] = {k: v["rows"] for k, v in collapsed.items()}

    # -- exposure index -------------------------------------------------------------
    units = {
        "credentials": summary["leaks"]["credentials"]["verified_login"] * 3 + (summary["leaks"]["credentials"]["total"] - summary["leaks"]["credentials"]["verified_login"]),
        "dark_web": len(dw_open),
        "code": sum(3 if c["severity"] == "high" else 1.5 if c["severity"] == "medium" else 0.5 for c in current(leaks["code"])),
        "files": summary["leaks"]["files"]["open"],
        "login_portals": len(cur_portals) + summary["risks"]["login_portals"]["plain_http"],
        "certificates": len(cert_risks),
        "vulnerabilities": len(vulnerabilities),
        "risky_services": 0,
        "malicious_ip_tags": 0,
    }
    exposure = exposure_index(units)

    # -- coverage ----------------------------------------------------------------------
    def cov(mod, status, count, note=None):
        return {"module": mod, "label": MODULE_LABELS[mod], "status": status, "count": count, "last_checked": args.report_date, "note": note}
    coverage = [
        cov("asset_discovery", "findings", summary["assets"]["domains"]["current"] + summary["assets"]["ips"]["total_rows"] + summary["assets"]["websites"]["total_rows"]),
        cov("mobile_apps", "findings", summary["assets"]["mobile_apps"]["total_rows"]),
        cov("social_accounts", "findings", summary["assets"]["wechat_official_accounts"]["total_rows"] + summary["assets"]["wechat_mini_programs"]["total_rows"]),
        cov("login_portals", "findings", len(cur_portals)),
        cov("risky_services", "monitored_none", 0, "No risky ports or services were identified this period."),
        cov("certificates", "findings" if cert_risks else "monitored_none", len(cert_risks)),
        cov("http_misconfigurations", "pending_semantics", len(assets["http_configs"]), "Header checks collected for every site; field semantics to be confirmed with the delivery team before customer display."),
        cov("malicious_ip_tags", "monitored_none", 0, "No IP addresses carried malicious reputation tags this period."),
        cov("vulnerabilities", "findings" if vulnerabilities else "monitored_none", len(vulnerabilities), None if vulnerabilities else "No exploitable vulnerabilities were identified this period."),
        cov("dark_web", "findings", len(dw_open)),
        cov("files", "findings", summary["leaks"]["files"]["open"]),
        cov("code", "findings", summary["leaks"]["code"]["current"]),
        cov("credentials", "findings", summary["leaks"]["credentials"]["total"]),
        cov("emails", "findings", summary["leaks"]["emails"]["total"]),
    ]

    # -- unified findings ----------------------------------------------------------------
    findings = []
    period_id = args.period_id

    def add_finding(module, item, title, entity_ref, severity, analyst_status, extra=None):
        findings.append({
            "id": f"{module}:{item['id']}", "module": module, "severity": severity, "title": title, "entity_ref": entity_ref,
            "lifecycle": item.get("lifecycle", "active"), "first_seen_period": period_id if item.get("lifecycle") == "new" else None,
            "analyst_status": analyst_status, "customer_status": "open", "owner": None, "customer_note": None,
            **(extra or {}),
        })

    for p in cur_portals:
        add_finding("login_portals", p, f"Exposed login portal: {p.get('title') or p['host']}", {"type": "login_portal", "id": p["id"]}, p["severity"], "open")
    for c in cert_risks:
        add_finding("certificates", c, f"Legacy TLS ({', '.join(c.get('protocols') or [])}) on {c['host']}:{c['port']}", {"type": "certificate", "id": c["id"]}, c["severity"], "open")
    for v in vulnerabilities:
        add_finding("vulnerabilities", v, v["name"], {"type": "vulnerability", "id": v["id"]}, v["severity"], "open")
    for d in leaks["dark_web"]:
        st = "taken_down" if d.get("taken_down") else "invalid" if d["lifecycle"] == "closed" else "to_confirm" if (d.get("analyst_note") or "").lower().startswith("to be confirmed") else "confirmed" if (d.get("analyst_note") or "").lower().startswith("confirmed") else "open"
        add_finding("dark_web", d, f"Dark web listing: {d['title_clean'][:80]}", {"type": "dark_web", "id": d["id"]}, d["severity"], st)
    for f in leaks["files"]:
        add_finding("files", f, f"Sensitive file shared publicly: {f.get('file_name') or f.get('url')}", {"type": "file", "id": f["id"]}, f["severity"], "taken_down" if f["taken_down"] else "open")
    for c in leaks["code"]:
        add_finding("code", c, f"Code repository referencing the organisation: {c['repository']}", {"type": "code", "id": c["id"]}, c["severity"], "expired" if c["lifecycle"] == "closed" else "open")
    for c in leaks["credentials"]:
        add_finding("credentials", c, f"Leaked credential for {c['host']}: {c['username']}", {"type": "credential", "id": c["id"]}, c["severity"], "verified" if c["verified_login"] else "open")
    # emails: one aggregated finding per period (1,000+ rows would flood the tracker)
    if leaks["emails"]:
        findings.append({"id": f"emails:{period_id}", "module": "emails", "severity": "low", "title": f"{len(leaks['emails'])} corporate email addresses exposed on the internet ({summary['leaks']['emails']['new']} new this period)",
                         "entity_ref": {"type": "email_set", "id": period_id}, "lifecycle": "active", "first_seen_period": None, "analyst_status": "open", "customer_status": "open", "owner": None, "customer_note": None, "aggregate_count": len(leaks["emails"])})

    # -- cross-checks vs report headings ----------------------------------------------------
    checks = []
    def chk(label, ours, theirs, collapsed_rows=0):
        if theirs is None:
            return
        item = {"item": label, "datapack": ours, "report": theirs, "match": ours == theirs}
        if not item["match"] and collapsed_rows and ours + collapsed_rows == theirs:
            item["match"] = True
            item["collapsed_rows"] = collapsed_rows
            item["note"] = f"report counted {collapsed_rows} duplicate row(s) that the datapack collapsed into existing entities"
        checks.append(item)
    chk("login portals (current)", len(cur_portals), section_counts.get("login_portals"), collapsed.get("login_portals", {}).get("current", 0))
    chk("certificate risks", len(cert_risks), section_counts.get("certificates"))
    chk("dark web items (open)", len(dw_open), section_counts.get("dark_web"))
    chk("file leaks (open)", summary["leaks"]["files"]["open"], section_counts.get("files"))
    chk("code leaks (current)", summary["leaks"]["code"]["current"], section_counts.get("code"))
    chk("credentials (non-empty rows)", summary["leaks"]["credentials"]["total"], section_counts.get("credentials"), collapsed.get("credentials", {}).get("rows", 0))
    chk("wechat official accounts", summary["assets"]["wechat_official_accounts"]["total_rows"], section_counts.get("social_accounts"))
    chk("mobile apps", summary["assets"]["mobile_apps"]["current"], section_counts.get("mobile_apps"), collapsed.get("mobile_apps", {}).get("current", 0))
    for kf in narrative["key_findings"]:
        if kf["module"] == "emails" and kf["count"] is not None:
            chk("emails (key findings)", summary["leaks"]["emails"]["total"], kf["count"])
        if kf["module"] == "asset_discovery" and kf["count"] is not None and "subdomain" in kf["text"].lower():
            chk("subdomains (key findings)", summary["assets"]["domains"]["current"], kf["count"])
    for c in checks:
        if not c["match"]:
            hint = ""
            if c["item"].startswith("credentials") and c["report"] and c["report"] - c["datapack"] == 1:
                hint = " (the sheet carries one blank row; the report likely counted it)"
            warnings.append(f"cross-check mismatch: {c['item']} datapack={c['datapack']} report={c['report']}{hint}")

    if getattr(args, "rebrand", ""):
        src_name, dst_name = args.rebrand.split(":", 1)
        def _rb(x):
            if isinstance(x, str):
                return x.replace(src_name, dst_name)
            if isinstance(x, list):
                return [_rb(i) for i in x]
            if isinstance(x, dict):
                return {k: _rb(v) for k, v in x.items()}
            return x
        narrative = _rb(narrative)
        warnings.append(f"narrative text rebranded: {src_name!r} -> {dst_name!r} (report-derived prose only)")

    previous = None
    if args.previous_summary:
        previous = json.loads(Path(args.previous_summary).read_text(encoding="utf-8"))

    pack = OrderedDict([
        ("schema_version", SCHEMA_VERSION),
        ("generator", {"name": "easm_convert.py", "version": GENERATOR_VERSION, "generated_at": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")}),
        ("customer", {"id": args.customer_id, "name": args.customer_name}),
        ("period", {"id": period_id, "label": args.period_label or period_id, "report_no": args.report_no, "report_date": args.report_date,
                    "previous_period_id": args.previous_period_id, "status": "draft"}),
        ("sources", sources),
        ("scope", {"root_domains": sorted({d["root_domain"] for d in assets["domains"] if d.get("root_domain")}), "scope_text": narrative["scope_text"]}),
        ("summary", summary),
        ("previous_summary", previous),
        ("exposure_index", exposure),
        ("narrative", {k: v for k, v in narrative.items() if k != "scope_text"}),
        ("coverage", coverage),
        ("assets", assets),
        ("risks", risks),
        ("leaks", leaks),
        ("findings", findings),
        ("cross_checks", checks),
        ("warnings", warnings),
    ])
    pack["_images"] = {"count": images.count, "bytes_in": images.bytes_in, "bytes_out": images.bytes_out}
    return pack


def write_report(pack: dict, out_dir: Path) -> None:
    s = pack["summary"]
    lines = [f"# Convert report — {pack['customer']['name']} — {pack['period']['label']}", "",
             f"Generated {pack['generator']['generated_at']} by {pack['generator']['name']} {pack['generator']['version']}", "",
             "## Counts", "",
             f"- Domains: {s['assets']['domains']['current']} current / {s['assets']['domains']['total_rows']} rows {s['assets']['domains']['by_lifecycle']}",
             f"- IPs: {s['assets']['ips']['total_rows']} rows, {s['assets']['ips']['current']} current, {s['assets']['ips']['with_open_ports']} with open ports",
             f"- Websites: {s['assets']['websites']['total_rows']}; Services: {s['assets']['services']['current']} current rows on {s['assets']['services']['ips_with_current_services']} IPs; Components: {s['assets']['components']['total_rows']} (admin panels {s['assets']['components']['admin_panels']})",
             f"- Certificates: {s['assets']['certificates']['total_rows']} rows; hygiene {s['assets']['certificates']['hygiene']}",
             f"- HTTP configs: {s['assets']['http_configs']['total_rows']} sites (semantics pending)",
             f"- Mobile apps: {s['assets']['mobile_apps']['current']} current / {s['assets']['mobile_apps']['total_rows']} rows; WeChat OA {s['assets']['wechat_official_accounts']['total_rows']}; mini programs {s['assets']['wechat_mini_programs']['total_rows']}",
             f"- Login portals: {s['risks']['login_portals']}",
             f"- Certificate risks: {s['risks']['certificates']}",
             f"- Dark web: {s['leaks']['dark_web']}",
             f"- Files: {s['leaks']['files']}",
             f"- Code: {s['leaks']['code']}",
             f"- Credentials: {s['leaks']['credentials']}",
             f"- Emails: {s['leaks']['emails']}",
             f"- Findings: {len(pack['findings'])}; Exposure index: {pack['exposure_index']['score']} ({pack['exposure_index']['band']}, grade {pack['exposure_index']['grade']})",
             f"- Screenshots: {pack['_images']['count']} files, {pack['_images']['bytes_in']/1e6:.1f} MB in -> {pack['_images']['bytes_out']/1e6:.1f} MB out",
             "", "## Cross-checks against the report", ""]
    for c in pack["cross_checks"]:
        lines.append(f"- {'OK ' if c['match'] else 'MISMATCH'} {c['item']}: datapack={c['datapack']} report={c['report']}")
    lines += ["", "## Warnings", ""] + ([f"- {w}" for w in pack["warnings"]] or ["- none"])
    (out_dir / "convert-report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--inventory", required=True)
    ap.add_argument("--report")
    ap.add_argument("--verification")
    ap.add_argument("--previous-summary")
    ap.add_argument("--out", required=True)
    ap.add_argument("--customer-id", default="customer")
    ap.add_argument("--customer-name", default="Customer")
    ap.add_argument("--period-id", required=True, help="e.g. 2026Q2")
    ap.add_argument("--period-label", default=None)
    ap.add_argument("--report-no", type=int, default=None)
    ap.add_argument("--report-date", required=True, help="YYYY-MM-DD (date of the report / scan reference)")
    ap.add_argument("--previous-period-id", default=None)
    ap.add_argument("--secret-key-env", default="EASM_SECRET_KEY")
    ap.add_argument("--max-image-width", type=int, default=1200)
    ap.add_argument("--rebrand", default="", help="Replace a legacy brand name in narrative text, e.g. 'OldBrand:NewBrand' (data fields are never touched)")
    args = ap.parse_args(argv)

    pack = build(args)
    out_dir = Path(args.out)
    (out_dir / "datapack.json").write_text(json.dumps(pack, ensure_ascii=False, indent=1), encoding="utf-8")
    write_report(pack, out_dir)
    print(f"datapack written: {out_dir / 'datapack.json'} ({(out_dir / 'datapack.json').stat().st_size/1e6:.1f} MB)")
    print(f"findings={len(pack['findings'])} exposure={pack['exposure_index']['score']} warnings={len(pack['warnings'])}")
    for w in pack["warnings"]:
        print("  WARN:", w)
    return 0


if __name__ == "__main__":
    sys.exit(main())
